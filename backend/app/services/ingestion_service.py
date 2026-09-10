import io
import re
import uuid
from pathlib import Path
from typing import List, Optional

import docx
from app.models.document_models import PageContent, ParsedDocument
from pypdf import PdfReader


class FileValidationError(Exception):
    """Raised when an uploaded file fails validation."""

    pass


class DecompressionBombError(Exception):
    """Raised when extracted document text exceeds safety limits."""

    pass


class IngestionService:
    """Service for validating, parsing, and sanitizing documents."""

    SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}
    DEFAULT_MAX_FILE_SIZE_MB = 15
    DEFAULT_MAX_CHARACTERS = 500_000

    def __init__(
        self,
        max_file_size_mb: int = DEFAULT_MAX_FILE_SIZE_MB,
        max_character_count: int = DEFAULT_MAX_CHARACTERS,
    ):
        self.max_file_size_bytes = max_file_size_mb * 1024 * 1024
        self.max_character_count = max_character_count

    def sanitize_text(self, text: str) -> str:
        """Sanitizes text by stripping null bytes and normalizing whitespace."""
        if not text:
            return ""

        # Replace null bytes & control chars (except \n, \r, \t) with space
        sanitized = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", text)

        # Normalize unicode non-breaking spaces and line separators
        sanitized = (
            sanitized.replace("\xa0", " ")
            .replace("\u2028", "\n")
            .replace("\u2029", "\n")
        )

        # Normalize excessive blank spaces per line
        lines = [
            re.sub(r"[ \t]+", " ", line).strip() for line in sanitized.splitlines()
        ]

        # Join with single newlines, preserving paragraph breaks
        result = []
        consecutive_empty = 0
        for line in lines:
            if not line:
                consecutive_empty += 1
                if consecutive_empty <= 1:
                    result.append("")
            else:
                consecutive_empty = 0
                result.append(line)

        return "\n".join(result).strip()

    def validate_file_metadata(self, filename: str, content: bytes) -> str:
        """Validates file extension, size limit, and magic byte headers."""
        if not content:
            raise FileValidationError("File is empty.")

        if len(content) > self.max_file_size_bytes:
            max_mb = self.max_file_size_bytes // (1024 * 1024)
            raise FileValidationError(
                f"File size exceeds the maximum allowed limit of {max_mb} MB."
            )

        ext = Path(filename).suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            allowed = ", ".join(sorted(self.SUPPORTED_EXTENSIONS))
            raise FileValidationError(
                f"Unsupported file extension '{ext}'. Supported formats: {allowed}"
            )

        # Magic byte header inspection
        if ext == ".pdf":
            if not content.startswith(b"%PDF-"):
                raise FileValidationError("Corrupted or invalid PDF header.")
        elif ext == ".docx":
            if not content.startswith(b"PK\x03\x04"):
                raise FileValidationError("Corrupted or invalid DOCX archive header.")
        elif ext in {".txt", ".md"}:
            # Ensure it is not an executable binary
            if content.startswith(b"MZ") or content.startswith(b"\x7fELF"):
                raise FileValidationError("Executable binary content is prohibited.")

        return ext

    def parse_pdf(self, content: bytes) -> List[PageContent]:
        """Extracts text page-by-page from a PDF byte stream."""
        pages: List[PageContent] = []
        try:
            reader = PdfReader(io.BytesIO(content))
            for idx, page in enumerate(reader.pages, start=1):
                raw_page_text = page.extract_text() or ""
                sanitized_page = self.sanitize_text(raw_page_text)
                pages.append(PageContent(page_number=idx, text=sanitized_page))
        except Exception as e:
            if isinstance(e, FileValidationError):
                raise
            raise FileValidationError(f"Failed to parse PDF document: {str(e)}") from e

        if not pages:
            pages.append(PageContent(page_number=1, text=""))
        return pages

    def parse_docx(self, content: bytes) -> List[PageContent]:
        """Extracts text from a DOCX document."""
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)
            sanitized = self.sanitize_text(full_text)
            return [PageContent(page_number=None, text=sanitized)]
        except Exception as e:
            raise FileValidationError(f"Failed to parse DOCX document: {str(e)}") from e

    def parse_plain_text(self, content: bytes) -> List[PageContent]:
        """Extracts text from plain text or markdown files."""
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception as e:
                raise FileValidationError(
                    f"Failed to decode text file: {str(e)}"
                ) from e

        sanitized = self.sanitize_text(text)
        return [PageContent(page_number=None, text=sanitized)]

    def validate_and_parse(
        self,
        filename: str,
        content: bytes,
        document_id: Optional[str] = None,
    ) -> ParsedDocument:
        """Main entry point to validate, parse, and enforce safety guards."""
        ext = self.validate_file_metadata(filename, content)

        if ext == ".pdf":
            pages = self.parse_pdf(content)
        elif ext == ".docx":
            pages = self.parse_docx(content)
        elif ext in {".txt", ".md"}:
            pages = self.parse_plain_text(content)
        else:
            raise FileValidationError(f"Unsupported format: {ext}")

        raw_combined_text = "\n\n".join(p.text for p in pages if p.text)

        # Decompression bomb guard
        if len(raw_combined_text) > self.max_character_count:
            raise DecompressionBombError(
                f"Text length exceeds limit of {self.max_character_count} characters."
            )

        doc_id = document_id or str(uuid.uuid4())
        return ParsedDocument(
            document_id=doc_id,
            filename=filename,
            file_size_bytes=len(content),
            pages=pages,
            raw_text=raw_combined_text,
        )
